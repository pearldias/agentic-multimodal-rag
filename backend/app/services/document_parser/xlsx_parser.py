"""Excel (.xlsx) document parser using openpyxl for legal schedules and tabular data."""

from datetime import datetime, timezone
from pathlib import Path
import openpyxl
from backend.app.core.exceptions import CorruptedDocumentError, EmptyDocumentError
from backend.app.models.document import DocumentMetadata, DocumentType, ParsedDocument, ParsedPage
from backend.app.services.document_parser.base import BaseDocumentParser
from backend.app.services.document_parser.normalizer import (
    format_table_as_markdown,
    normalize_legal_text,
)


class XlsxParser(BaseDocumentParser):
    """Parser for XLSX spreadsheets using openpyxl, preserving sheets and table structures."""

    @property
    def supported_extensions(self) -> set[str]:
        return {"xlsx"}

    def parse(
        self,
        file_path: Path,
        source_url: str | None = None,
        title: str | None = None,
        **kwargs,
    ) -> ParsedDocument:
        path = Path(file_path)
        if not path.is_file():
            raise FileNotFoundError(f"File not found: {path}")

        file_size = path.stat().st_size
        if file_size == 0:
            raise EmptyDocumentError(path.name, reason="File is 0 bytes")

        try:
            # data_only=True returns evaluated values instead of raw formulas
            wb = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
        except Exception as exc:
            raise CorruptedDocumentError(path.name, details=str(exc)) from exc

        try:
            pages: list[ParsedPage] = []
            consolidated_texts: list[str] = []

            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                rows_data: list[list[str]] = []

                for row in ws.iter_rows(values_only=True):
                    # Check if row has any non-empty cell
                    if any(cell is not None and str(cell).strip() for cell in row):
                        cleaned_row = [
                            str(cell).strip().replace("\n", " ") if cell is not None else ""
                            for cell in row
                        ]
                        rows_data.append(cleaned_row)

                if not rows_data:
                    continue  # Skip completely empty sheets

                headers = rows_data[0]
                body_rows = rows_data[1:] if len(rows_data) > 1 else []
                table_md = format_table_as_markdown(headers, body_rows)

                sheet_content = f"### Sheet: {sheet_name}\n\n{table_md}"
                normalized_sheet_content = normalize_legal_text(sheet_content)

                if normalized_sheet_content:
                    pages.append(
                        ParsedPage(
                            text=normalized_sheet_content,
                            sheet_name=sheet_name,
                            metadata={
                                "filename": path.name,
                                "sheet_name": sheet_name,
                                "row_count": len(rows_data),
                            },
                        )
                    )
                    consolidated_texts.append(normalized_sheet_content)

            if not pages or not consolidated_texts:
                raise EmptyDocumentError(
                    path.name,
                    reason="Excel workbook contains no non-empty rows or readable data",
                )

            resolved_title = title or path.stem.replace("_", " ").title()
            timestamp = datetime.now(timezone.utc).isoformat()

            metadata = DocumentMetadata(
                filename=path.name,
                file_type=DocumentType.XLSX.value,
                title=resolved_title,
                source_url=source_url,
                ingestion_timestamp=timestamp,
                file_size_bytes=file_size,
                total_pages=len(pages),
                extra={"sheet_names": wb.sheetnames},
            )

            raw_text = "\n\n".join(consolidated_texts)

            return ParsedDocument(
                metadata=metadata,
                pages=pages,
                raw_text=raw_text,
            )
        finally:
            wb.close()
